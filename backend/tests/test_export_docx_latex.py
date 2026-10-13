"""导出链路回归测试：DOCX 转换、LaTeX 生成、中文文件名响应头。

背景（2026-10-13 全量功能测试实测发现三个真缺陷）：
1. `POST /api/export/docx` 对任何中文标题都 500 —— HTTP 头按 latin-1 编码，装不下中文；
2. `?format=latex` 导出的其实是 Python 代码 —— 代码块正则把语言标记写成可选（````(?:latex|tex)?````）；
3. 任务导出的 Word 表格全丢、``**加粗**`` 原样显示 —— `api/tasks.py` 里另有一份简化实现。
"""

import docx
import pytest

from app.services.md_docx import (
    content_disposition,
    extract_explicit_latex,
    markdown_to_docx,
    markdown_to_latex,
)

SAMPLE = """# 一、问题重述

某车间用两种原料生产三种产品，目标是 **利润最大化**，模型如下：

$$ \\max Z = 5x_1 + 4x_2 + 3x_3 $$

行内公式 $x_1 \\ge 0$ 与特殊符号 & 百分号 % 下划线 _ 都要安全。

| 产品 | 利润 | 原料A |
| --- | --- | --- |
| 甲 | 5 | 2 |
| 乙 | 4 | 1 |

- 列表项一
- 列表项二

```python
print("代码块内容必须保留")
```
"""


# ── 1. 响应头：中文文件名 ──────────────────────────────────────────


def test_content_disposition_survives_latin1_encoding() -> None:
    """核心回归：中文文件名必须能按 latin-1 编码（否则 Starlette 抛 UnicodeEncodeError → 500）。"""
    header = content_disposition("数学建模方案.docx", fallback="export.docx")
    header.encode("latin-1")
    assert "filename*=UTF-8''" in header
    assert "%E6%95%B0%E5%AD%A6" in header  # “数学”的 percent-encoding


def test_content_disposition_ascii_name_kept_as_fallback() -> None:
    header = content_disposition("paper.docx")
    assert 'filename="paper.docx"' in header
    assert "filename*=UTF-8''paper.docx" in header


def test_content_disposition_all_non_ascii_keeps_extension() -> None:
    """全非 ASCII 文件名：ASCII 回退名只保留扩展名，真实名字走 filename*。"""
    header = content_disposition("论文.docx", fallback="export.docx")
    header.encode("latin-1")
    assert "filename*=UTF-8''%E8%AE%BA%E6%96%87.docx" in header
    assert ".docx" in header


# ── 2. DOCX：表格与行内格式 ────────────────────────────────────────


def test_docx_keeps_tables() -> None:
    """回归：旧的任务导出实现直接 continue 掉所有 `|` 行，Word 里表格为 0。"""
    doc = docx.Document()
    markdown_to_docx(doc, SAMPLE)
    assert len(doc.tables) == 1
    assert doc.tables[0].cell(1, 0).text == "甲"
    assert doc.tables[0].cell(0, 1).text == "利润"


def test_docx_parses_inline_formatting() -> None:
    doc = docx.Document()
    markdown_to_docx(doc, SAMPLE)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "**" not in text, "加粗标记必须被解析掉"
    assert "利润最大化" in text
    # 行内公式降级为可读文本；而独立公式（$$...$$）已升级为 Word 原生公式，
    # 因此不会再以 [公式: ...] 的形式出现在纯文本段落里。
    assert "[x_1 " in text, "行内公式应降级为可读文本"


def test_docx_keeps_code_block_content() -> None:
    doc = docx.Document()
    markdown_to_docx(doc, SAMPLE)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "代码块内容必须保留" in text


def test_docx_headings_and_lists() -> None:
    doc = docx.Document()
    markdown_to_docx(doc, SAMPLE)
    style_names = {p.style.name for p in doc.paragraphs}
    assert "Heading 1" in style_names
    texts = [p.text for p in doc.paragraphs]
    assert any(t.startswith("• ") for t in texts), "无序列表应转成带项目符号的段落"


# ── 3. LaTeX：提取与生成 ───────────────────────────────────────────


def test_extract_explicit_latex_ignores_other_languages() -> None:
    """回归：语言标记可选的正则会抓走任意首个代码块（实测导出的是 Python 代码）。"""
    text = "```python\nprint(1)\n```\n\n```latex\n\\documentclass{ctexart}\n```"
    got = extract_explicit_latex(text)
    assert got is not None
    assert got.startswith("\\documentclass")


def test_extract_explicit_latex_returns_none_without_explicit_block() -> None:
    assert extract_explicit_latex("```python\nprint(1)\n```") is None
    assert extract_explicit_latex("普通正文，没有代码块") is None
    assert extract_explicit_latex("```\n无语言标记\n```") is None


def test_markdown_to_latex_skeleton_is_compilable() -> None:
    tex = markdown_to_latex(SAMPLE, title="优化模型")
    assert tex.startswith("\\documentclass")
    assert "ctexart" in tex
    assert "\\begin{document}" in tex and "\\end{document}" in tex
    assert "\\section{一、问题重述}" in tex


def test_markdown_to_latex_keeps_math_and_table() -> None:
    tex = markdown_to_latex(SAMPLE)
    assert "\\max Z = 5x_1" in tex, "行间公式应原样保留"
    assert "\\begin{tabular}" in tex and "甲" in tex


def test_markdown_to_latex_escapes_specials_outside_math() -> None:
    tex = markdown_to_latex("普通文本 & 特殊 % 字符 _ 下划线 与公式 $a_1 + b_2$")
    body = tex.split("\\begin{document}", 1)[1]
    assert "\\&" in body and "\\%" in body
    assert "$a_1 + b_2$" in body, "公式内部不得被转义"


def test_markdown_to_latex_python_block_becomes_verbatim() -> None:
    tex = markdown_to_latex("# 标题\n\n```python\nimport numpy\n```")
    assert "\\begin{verbatim}" in tex
    assert "import numpy" in tex
    assert tex.startswith("\\documentclass")


# ── 4. 公式：Word 原生公式（OMML） ─────────────────────────────────


def _omml_available() -> bool:
    pytest.importorskip("latex2mathml")
    from app.services.md_docx import _find_mml2omml_xsl

    return _find_mml2omml_xsl() is not None


def test_latex_to_omml_produces_math_element() -> None:
    """LaTeX → OMML 应产出 Word 原生公式节点（缺 Office XSL 时跳过）。"""
    if not _omml_available():
        pytest.skip("本机无 MML2OMML.XSL（未安装 Office）")
    from app.services.md_docx import latex_to_omml

    omml = latex_to_omml(r"\frac{a}{b}")
    assert omml is not None
    assert "<m:oMath" in omml


def test_docx_contains_native_formula_not_text(tmp_path) -> None:
    """回归：独立公式块必须写进 docx 作为 <m:oMath>，不再是 [公式: ...] 文本降级。"""
    if not _omml_available():
        pytest.skip("本机无 MML2OMML.XSL（未安装 Office）")
    import zipfile

    doc = docx.Document()
    markdown_to_docx(doc, "$$\\max\\ Z = 5x_1 + 4x_2$$\n\n| a | b |\n| --- | --- |\n| 1 | 2 |")
    path = tmp_path / "formula.docx"
    doc.save(str(path))

    xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8")
    assert "<m:oMath" in xml, "独立公式应转成 Word 原生公式"
    assert "[公式: " not in xml, "不应残留文本降级标记"
    assert len(doc.tables) == 1, "公式处理不得影响表格"


def test_formula_falls_back_to_text_without_xslt(monkeypatch) -> None:
    """拿不到 XSLT 时必须平滑降级为文本、不得抛错（未装 Office 的机器仍要能导出）。"""
    from app.services import md_docx

    monkeypatch.setattr(md_docx, "_get_xslt", lambda: None)
    md_docx.latex_to_omml.cache_clear()
    try:
        doc = docx.Document()
        md_docx.markdown_to_docx(doc, "$$\\frac{a}{b}$$")
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "[公式: " in text
    finally:
        md_docx.latex_to_omml.cache_clear()


def test_docx_multiline_formula_block(tmp_path) -> None:
    """回归：多行公式块（`$$` 独占一行包裹 `\\begin{cases}`）在真实论文里约占多数
    ——实测一篇 14 个 vs 单行块 1 个，必须与单行块一样转成原生公式。"""
    if not _omml_available():
        pytest.skip("本机无 MML2OMML.XSL（未安装 Office）")
    import zipfile

    md = (
        "$$\n"
        "\\begin{cases}\n"
        "x_1 = t + 20,\\\\\n"
        "x_2 = t,\n"
        "\\end{cases}\n"
        "$$\n\n"
        "正文段落\n"
    )
    doc = docx.Document()
    markdown_to_docx(doc, md)
    path = tmp_path / "multiline.docx"
    doc.save(str(path))

    xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8")
    assert "<m:oMath" in xml, "多行公式块应转成 Word 原生公式"
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "$$" not in text, "公式定界符不得作为正文残留"
