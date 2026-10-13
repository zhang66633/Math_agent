"""Markdown → DOCX / LaTeX 的共享转换实现（单一真源）。

**为什么要有这个模块**：此前存在两份 DOCX 转换——

- `api/export_routes.py` 的版本解析完整（标题 / 列表 / 表格 / 引用 / 代码块 / 行内格式 / 公式降级）；
- `api/tasks.py` 里 `_export_docx` 的版本会**跳过表格**、也不解析行内标记。

同一篇论文走两个入口导出，结果不一致（实测任务导出的 Word 里表格为 0、`**加粗**` 原样显示）。
现在统一到本模块，两个入口共用；图片由调用方通过 `image_resolver` 注入（只有任务导出需要插图）。

另含两个下载侧工具：
- `content_disposition()`：HTTP 头只能装 latin-1，中文文件名直接塞进去会抛
  `UnicodeEncodeError`（实测 `/api/export/docx` 对任何中文标题都 500），按 RFC 5987 修复；
- `markdown_to_latex()`：把 markdown 正文转成可编译的 ctexart 文档。
"""

from __future__ import annotations

import logging
import re
from collections.abc import Callable
from functools import lru_cache
from urllib.parse import quote

logger = logging.getLogger(__name__)

# ══════════════════════════════════════════════════════════════════
# 下载响应头：中文文件名
# ══════════════════════════════════════════════════════════════════


def content_disposition(filename: str, fallback: str = "download") -> str:
    """生成能承载中文文件名的 Content-Disposition（RFC 5987）。

    Starlette 会把头值按 latin-1 编码，中文名会直接抛 UnicodeEncodeError；
    因此同时给出 ASCII 回退名（老浏览器）与 `filename*=UTF-8''` 编码名（现代浏览器）。
    """
    ascii_fallback = "".join(
        ch if (ch.isascii() and ch not in '"\\') else "_" for ch in filename
    ).strip()
    if not ascii_fallback or set(ascii_fallback) <= {"_", "."}:
        ascii_fallback = fallback
    return f"attachment; filename=\"{ascii_fallback}\"; filename*=UTF-8''{quote(filename, safe='')}"


# ══════════════════════════════════════════════════════════════════
# 行内 Markdown → 纯文本
# ══════════════════════════════════════════════════════════════════


def parse_inline(text: str) -> str:
    """把一行内的 Markdown 标记转成纯文本（与 marked 渲染后的可读内容对齐）。"""
    # 转义字符先占位保护（\* \_ \`），避免被下方格式正则误吞，最后还原
    text = text.replace(r"\*", "\x00A").replace(r"\_", "\x00B").replace(r"\`", "\x00C")
    # 行内代码
    text = re.sub(r"`([^`]+)`", r"\1", text)
    # 图片 ![alt](url) → alt（图片: url）
    text = re.sub(r"!\[([^\]]*)\]\(([^)]+)\)", r"\1（图片: \2）", text)
    # 链接 [text](url) → text（url）
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1（\2）", text)
    # 加粗（** / __）
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"__(.+?)__", r"\1", text)
    # 斜体（剩余单星）
    text = re.sub(r"\*([^*\n]+)\*", r"\1", text)
    # 删除线
    text = re.sub(r"~~(.+?)~~", r"\1", text)
    # 公式：$$..$$ 与 $..$
    text = re.sub(r"\$\$(.+?)\$\$", r"[公式: \1]", text)
    text = re.sub(r"\$(.+?)\$", r"[\1]", text)
    # 还原转义字符
    text = text.replace("\x00A", "*").replace("\x00B", "_").replace("\x00C", "`")
    return text.strip()


# 保持与历史公开名一致（测试与旧调用点使用下划线前缀）
_parse_inline = parse_inline


# ══════════════════════════════════════════════════════════════════
# 公式：LaTeX → MathML → OMML（Word 原生公式）
# ══════════════════════════════════════════════════════════════════

# MML2OMML.XSL 是 Office 自带的 MathML→OMML 转换表。用它可以把公式写成 Word
# **原生公式**（可编辑、可编号、随文排版），而不是 LaTeX 源码文本。
# 找不到该文件或任一环节失败时，公式回落为文本降级，绝不阻断导出。
_XSL_CANDIDATES = (
    r"C:\Program Files\Microsoft Office\root\OfficeXX\MML2OMML.XSL",
    r"C:\Program Files (x86)\Microsoft Office\root\OfficeXX\MML2OMML.XSL",
    r"C:\Program Files\Microsoft Office\OfficeXX\MML2OMML.XSL",
)


def _find_mml2omml_xsl() -> str | None:
    """定位 MML2OMML.XSL：优先环境变量 MML2OMML_XSL，其次 Office 常见安装路径。"""
    import os
    from pathlib import Path as _Path

    env = os.getenv("MML2OMML_XSL")
    if env and _Path(env).exists():
        return env
    for pattern in _XSL_CANDIDATES:
        for ver in ("16", "15", ""):
            candidate = pattern.replace("XX", ver)
            if _Path(candidate).exists():
                return candidate
    for base in (
        r"C:\Program Files\Microsoft Office",
        r"C:\Program Files (x86)\Microsoft Office",
    ):
        root = _Path(base)
        if root.exists():
            for found in root.rglob("MML2OMML.XSL"):
                return str(found)
    return None


_XSLT_CACHE: dict = {"transform": None, "tried": False}


def _get_xslt():
    """惰性加载并缓存 XSLT 变换对象（XSL 近 200 KB，不能每个公式重新解析）。"""
    if _XSLT_CACHE["tried"]:
        return _XSLT_CACHE["transform"]
    _XSLT_CACHE["tried"] = True
    path = _find_mml2omml_xsl()
    if not path:
        logger.info("未找到 MML2OMML.XSL，Word 导出的公式将回落为文本")
        return None
    try:
        from lxml import etree

        _XSLT_CACHE["transform"] = etree.XSLT(etree.parse(path))
    except Exception as exc:  # noqa: BLE001
        logger.warning("加载 MML2OMML.XSL 失败（公式回落为文本）: %s", exc)
        _XSLT_CACHE["transform"] = None
    return _XSLT_CACHE["transform"]


@lru_cache(maxsize=512)
def latex_to_omml(latex: str) -> str | None:
    """LaTeX → OMML（Word 原生公式）。任一环节不可用返回 None，由调用方降级。

    链路：`latex2mathml` 转 MathML → Office 的 MML2OMML.XSL 转 OMML。
    两个依赖都是可选的，缺失时静默降级。
    """
    text = (latex or "").strip()
    if not text:
        return None
    transform = _get_xslt()
    if transform is None:
        return None
    try:
        import latex2mathml.converter as conv
        from lxml import etree

        mathml = conv.convert(text)
        omml = str(transform(etree.fromstring(mathml.encode("utf-8"))))
    except Exception as exc:  # noqa: BLE001 — 公式转换失败不该阻断导出
        logger.debug("LaTeX → OMML 失败（%s）: %s", text[:40], exc)
        return None
    for prefix in ('<?xml version="1.0"?>', "<?xml version='1.0'?>"):
        if omml.startswith(prefix):
            omml = omml[len(prefix) :]
    omml = omml.strip()
    return omml if "<m:oMath" in omml else None


def insert_omml_paragraph(doc, latex: str) -> bool:
    """把 LaTeX 公式作为**居中的 Word 原生公式**插入段落；不可用时返回 False。"""
    omml = latex_to_omml(latex)
    if not omml:
        return False
    try:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import parse_xml

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p._p.append(parse_xml(omml))
        return True
    except Exception as exc:  # noqa: BLE001
        logger.debug("OMML 注入 docx 失败: %s", exc)
        return False


# ══════════════════════════════════════════════════════════════════
# Markdown → DOCX
# ══════════════════════════════════════════════════════════════════

_LIST_ITEM_RE = re.compile(r"^\s*(?:[-*+]|\d+[.)])\s+(.*)$")
_TASK_ITEM_RE = re.compile(r"^\s*[-*+]\s+\[([ xX])\]\s+(.*)$")
_TABLE_ROW_RE = re.compile(r"^\s*\|.+\|\s*$")
_IMAGE_LINE_RE = re.compile(r"^!\[[^\]]*\]\((/api/(?:task_files|images)/[^)\s]+)\)$")


def is_table_separator(line: str) -> bool:
    """GFM 表格的分隔行（| --- | --- |）。"""
    return bool(re.match(r"^\s*\|[\s:|-]+\|\s*$", line))


_is_table_separator = is_table_separator


def _add_table(doc, rows: list[str]) -> None:
    """把 GFM 表格行列表转成 docx 表格（跳过分隔行）。"""
    parsed: list[list[str]] = []
    for row in rows:
        if is_table_separator(row):
            continue
        cells = [c.strip() for c in row.strip().strip("|").split("|")]
        parsed.append(cells)
    if not parsed:
        return
    n_cols = max(len(r) for r in parsed)
    table = doc.add_table(rows=len(parsed), cols=n_cols)
    table.style = "Table Grid"
    for ri, row in enumerate(parsed):
        for ci in range(n_cols):
            cell_text = parse_inline(row[ci]) if ci < len(row) else ""
            table.cell(ri, ci).text = cell_text


def _add_list_item(doc, line: str) -> None:
    """列表/任务列表 → 缩进段落。"""
    from docx.shared import Pt

    m = _TASK_ITEM_RE.match(line)
    if m:
        prefix = "☑ " if m.group(1).lower() == "x" else "☐ "
        text = f"{prefix}{m.group(2)}"
    else:
        m = _LIST_ITEM_RE.match(line)
        if not m:
            return
        numbered = bool(re.match(r"^\s*\d+[.)]\s", line))
        prefix = "1. " if numbered else "• "
        text = f"{prefix}{m.group(1)}"
    p = doc.add_paragraph(parse_inline(text))
    p.paragraph_format.left_indent = Pt(24)


def markdown_to_docx(
    doc,
    markdown: str,
    image_resolver: Callable[[str], object] | None = None,
) -> None:
    """逐块解析 Markdown 并写入 docx。

    Args:
        doc: python-docx ``Document``。
        markdown: 正文。
        image_resolver: 可选。传入后，独立的图片行会解析成本地路径并插入图片；
            返回 falsy 则回退成文本（`图片: url`）。
    """
    from docx.shared import Pt, RGBColor

    lines = markdown.split("\n")
    i = 0
    in_code_block = False
    code_lines: list[str] = []

    while i < len(lines):
        line = lines[i]

        # 代码块
        if line.strip().startswith("```"):
            if in_code_block:
                code_text = "\n".join(code_lines)
                p = doc.add_paragraph()
                run = p.add_run(code_text)
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # 独立图片行 → 插入本地图表（仅当调用方提供了解析器）
        if image_resolver is not None:
            img_match = _IMAGE_LINE_RE.match(line.strip())
            if img_match:
                try:
                    path = image_resolver(img_match.group(1))
                    if path:
                        from docx.enum.text import WD_ALIGN_PARAGRAPH
                        from docx.shared import Inches

                        doc.add_picture(str(path), width=Inches(5.5))
                        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        i += 1
                        continue
                except Exception:  # noqa: BLE001 — 插图失败不阻断导出
                    pass

        # 独立公式块 $$...$$ → Word 原生公式（OMML）；不可用时落到后面的文本降级
        formula_line = line.strip()
        if (
            formula_line.startswith("$$")
            and formula_line.endswith("$$")
            and len(formula_line) > 4
        ):
            if insert_omml_paragraph(doc, formula_line[2:-2].strip()):
                i += 1
                continue

        # 表格（连续 | 行聚合成一个 docx table）
        if _TABLE_ROW_RE.match(line):
            table_rows = []
            while i < len(lines) and _TABLE_ROW_RE.match(lines[i]):
                table_rows.append(lines[i])
                i += 1
            _add_table(doc, table_rows)
            continue

        # 标题
        if line.startswith("### "):
            doc.add_heading(parse_inline(line[4:]), level=3)
        elif line.startswith("## "):
            doc.add_heading(parse_inline(line[3:]), level=2)
        elif line.startswith("# "):
            doc.add_heading(parse_inline(line[2:]), level=1)
        # 分隔线
        elif line.strip() in ("---", "***", "___"):
            doc.add_paragraph("─" * 40)
        # 引用
        elif line.startswith("> "):
            p = doc.add_paragraph(parse_inline(line[2:]))
            p.paragraph_format.left_indent = Pt(24)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        # 列表 / 任务列表
        elif _LIST_ITEM_RE.match(line):
            _add_list_item(doc, line)
        # 空行
        elif not line.strip():
            pass
        # 普通段落（行内标记 → 纯文本）
        else:
            doc.add_paragraph(parse_inline(line))

        i += 1


_markdown_to_docx = markdown_to_docx


# ══════════════════════════════════════════════════════════════════
# Markdown → LaTeX（可编译的最简文档）
# ══════════════════════════════════════════════════════════════════

_LATEX_PREAMBLE = r"""\documentclass[12pt,a4paper]{ctexart}
\usepackage{amsmath,amssymb,graphicx,booktabs,geometry,longtable,hyperref}
\geometry{margin=2.5cm}
"""

_EXPLICIT_LATEX_BLOCK_RE = re.compile(r"```(?:latex|tex)\s*\n(.*?)```", re.DOTALL)


def extract_explicit_latex(text: str) -> str | None:
    """只提取**显式**标注为 latex/tex 的代码块。

    历史实现把语言标记写成可选的 ````(?:latex|tex)?````，于是会匹配**任意**首个代码块——
    实测导出的 .tex 是 Python 代码。这里要求必须显式标注。
    """
    m = _EXPLICIT_LATEX_BLOCK_RE.search(text or "")
    return m.group(1).strip() if m else None


def _escape_latex_text(text: str) -> str:
    """转义 LaTeX 特殊字符，但**保护行内公式** $...$ 不被转义。"""
    parts = re.split(r"(\$[^$\n]+\$)", text)
    out: list[str] = []
    for idx, part in enumerate(parts):
        if idx % 2 == 1:  # 公式段，原样
            out.append(part)
            continue
        s = part.replace("\\", "\x00B")  # 先保护反斜杠，避免后续替换互相干扰
        for ch, rep in (
            ("&", r"\&"),
            ("%", r"\%"),
            ("#", r"\#"),
            ("_", r"\_"),
            ("{", r"\{"),
            ("}", r"\}"),
            ("~", r"\textasciitilde{}"),
            ("^", r"\textasciicircum{}"),
        ):
            s = s.replace(ch, rep)
        out.append(s.replace("\x00B", r"\textbackslash{}"))
    return "".join(out)


def _latex_table(rows: list[str]) -> str:
    parsed: list[list[str]] = []
    for row in rows:
        if is_table_separator(row):
            continue
        parsed.append([c.strip() for c in row.strip().strip("|").split("|")])
    if not parsed:
        return ""
    n_cols = max(len(r) for r in parsed)
    spec = "|" + "l|" * n_cols
    lines = [r"\begin{center}", r"\begin{tabular}{" + spec + "}", r"\hline"]
    for ri, row in enumerate(parsed):
        cells = []
        for c in row:
            esc = _escape_latex_text(c)
            cells.append(r"\textbf{" + esc + "}" if ri == 0 else esc)
        cells += [""] * (n_cols - len(cells))
        lines.append(" & ".join(cells) + r" \\ \hline")
    lines += [r"\end{tabular}", r"\end{center}"]
    return "\n".join(lines)


def markdown_to_latex(markdown: str, title: str = "数学建模论文") -> str:
    """把 Markdown 正文转成可直接编译的 ctexart 文档。

    只保证两件事：**内容不丢**、**能编译**。标题映射章节层级，表格转 tabular，
    公式（行内 `$...$` 与行间 `$$...$$`）原样保留，代码块进 verbatim，
    其余文本按 LaTeX 规则转义。
    """
    body: list[str] = []
    lines = (markdown or "").split("\n")
    i = 0
    in_code = False
    in_list: str | None = None  # "itemize" | "enumerate"

    def close_list() -> None:
        nonlocal in_list
        if in_list:
            body.append(f"\\end{{{in_list}}}")
            in_list = None

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            close_list()
            if in_code:
                body.append(r"\end{verbatim}")
                in_code = False
            else:
                body.append(r"\begin{verbatim}")
                in_code = True
            i += 1
            continue
        if in_code:
            body.append(line)
            i += 1
            continue

        if _TABLE_ROW_RE.match(line):
            close_list()
            rows = []
            while i < len(lines) and _TABLE_ROW_RE.match(lines[i]):
                rows.append(lines[i])
                i += 1
            body.append(_latex_table(rows))
            continue

        if stripped.startswith("### "):
            close_list()
            body.append(r"\subsubsection{" + _escape_latex_text(stripped[4:]) + "}")
        elif stripped.startswith("## "):
            close_list()
            body.append(r"\subsection{" + _escape_latex_text(stripped[3:]) + "}")
        elif stripped.startswith("# "):
            close_list()
            body.append(r"\section{" + _escape_latex_text(stripped[2:]) + "}")
        elif stripped in ("---", "***", "___"):
            close_list()
            body.append(r"\par\noindent\rule{\linewidth}{0.4pt}")
        elif stripped.startswith("> "):
            close_list()
            body.append(r"\begin{quote}" + _escape_latex_text(stripped[2:]) + r"\end{quote}")
        elif _LIST_ITEM_RE.match(line):
            numbered = bool(re.match(r"^\s*\d+[.)]\s", line))
            want = "enumerate" if numbered else "itemize"
            if in_list != want:
                close_list()
                body.append(f"\\begin{{{want}}}")
                in_list = want
            item = _LIST_ITEM_RE.match(line).group(1)  # type: ignore[union-attr]
            body.append(r"\item " + _escape_latex_text(item))
        elif not stripped:
            close_list()
            body.append("")
        else:
            close_list()
            body.append(_escape_latex_text(line))
        i += 1

    close_list()
    if in_code:
        body.append(r"\end{verbatim}")

    return (
        _LATEX_PREAMBLE
        + r"\title{" + _escape_latex_text(title) + "}\n"
        + r"\begin{document}" + "\n"
        + r"\maketitle" + "\n"
        + "\n".join(body)
        + "\n" + r"\end{document}" + "\n"
    )
